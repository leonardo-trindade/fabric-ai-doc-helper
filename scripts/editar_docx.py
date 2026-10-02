"""Aplica alterações pontuais em um .docx JÁ REVISADO pelo usuário, sem regenerar o documento.

Preserva tudo que não for alterado (edições humanas, formatação, comentários). Sempre grava
um NOVO arquivo; nunca sobrescreve.

Uso:
    uv run python scripts/editar_docx.py --listar projeto/docs/revisado/DT_x_v0.2.docx [--filtro texto]
    uv run python scripts/editar_docx.py projeto/docs/alteracoes_v0.3.yaml

Especificação (YAML):
    origem: projeto/docs/revisado/DT_x_v0.2.docx
    saida: projeto/docs/DT_x_v0.3.docx
    alteracoes:
      - substituir: {de: "texto antigo", para: "texto novo"}            # todas as ocorrências (use max: 1 para limitar)
      - substituir_paragrafo: {contendo: "trecho único", texto: "novo texto do parágrafo"}
      - inserir_apos: {contendo: "trecho único", texto: "novo parágrafo (mesmo estilo do âncora)"}
      - remover_paragrafo: {contendo: "trecho único"}
      - historico: {versao: "0.3", data: "dd/mm/aaaa", autor: "Nome", descricao: "O que mudou"}

Regras: `contendo` precisa identificar UM parágrafo (texto de corpo ou de célula de tabela);
se não achar ou achar mais de um, a alteração falha e nada é gravado.
Depois: powershell -ExecutionPolicy Bypass -File scripts/finalizar_docx.ps1 <saida>  (sumário).
"""
from __future__ import annotations

import argparse
import copy
import sys
from pathlib import Path

import docx
import yaml
from docx.table import Table
from docx.text.paragraph import Paragraph

RAIZ = Path(__file__).resolve().parent.parent


def todos_paragrafos(d) -> list[Paragraph]:
    out: list[Paragraph] = []

    def de_bloco(bloco):
        for p in bloco.paragraphs:
            out.append(p)
        for t in bloco.tables:
            for row in t.rows:
                for cell in row.cells:
                    de_bloco(cell)

    de_bloco(d)
    for sec in d.sections:
        for parte in (sec.header, sec.footer):
            de_bloco(parte)
    # células mescladas aparecem repetidas: remove duplicatas pelo elemento XML
    vistos, unicos = set(), []
    for p in out:
        if id(p._p) not in vistos:
            vistos.add(id(p._p))
            unicos.append(p)
    return unicos


DOURADO = "C79408"  # cor dos marcadores de pendência {{...}} no template


def limpar_marcador(r) -> None:
    """Texto que deixou de ser pendência perde o estilo dourado/itálico do marcador."""
    if "{{" in r.text:
        return
    cor = r.font.color
    if cor is not None and cor.type is not None and str(cor.rgb).upper() == DOURADO:
        r.font.color.rgb = None
        r.font.italic = None


def trocar_texto(p: Paragraph, de: str, para: str, restante: int) -> tuple[int, bool]:
    feitas, achatado = _trocar(p, de, para, restante)
    for r in p.runs:
        if feitas and r.text:
            limpar_marcador(r)
    return feitas, achatado


def _trocar(p: Paragraph, de: str, para: str, restante: int) -> tuple[int, bool]:
    """Substitui dentro dos runs; se o trecho cruza runs, junta o parágrafo no 1º run."""
    feitas, achatado = 0, False
    for r in p.runs:
        while de in r.text and (restante < 0 or feitas < restante):
            r.text = r.text.replace(de, para, 1)
            feitas += 1
    if de in p.text and (restante < 0 or feitas < restante) and p.runs:
        n = p.text.count(de) if restante < 0 else min(p.text.count(de), restante - feitas)
        novo = p.text.replace(de, para, n)
        p.runs[0].text = novo
        for r in p.runs[1:]:
            r.text = ""
        feitas += n
        achatado = True
    return feitas, achatado


def unico(pars: list[Paragraph], trecho: str) -> Paragraph:
    achados = [p for p in pars if trecho in p.text]
    if len(achados) != 1:
        raise ValueError(f"'{trecho[:60]}' encontrado em {len(achados)} parágrafos (precisa ser exatamente 1).")
    return achados[0]


def definir_texto(p: Paragraph, texto: str) -> None:
    if p.runs:
        p.runs[0].text = texto
        for r in p.runs[1:]:
            r.text = ""
        limpar_marcador(p.runs[0])
    else:
        p.add_run(texto)


def aplicar(spec: dict) -> None:
    origem = RAIZ / spec["origem"]
    saida = RAIZ / spec["saida"]
    if saida.exists():
        sys.exit(f"{spec['saida']} já existe: use um novo nome/versão (nunca sobrescrever).")
    d = docx.Document(origem)
    avisos: list[str] = []
    for i, alt in enumerate(spec.get("alteracoes", []), 1):
        (op, v), = alt.items()
        pars = todos_paragrafos(d)
        if op == "substituir":
            total, restante = 0, int(v.get("max", -1))
            for p in pars:
                n, achatou = trocar_texto(p, v["de"], v["para"], -1 if restante < 0 else restante - total)
                total += n
                if achatou:
                    avisos.append(f"#{i}: trecho cruzava formatações; parágrafo '{p.text[:50]}…' ficou com a formatação do início.")
                if 0 <= restante <= total:
                    break
            if total == 0:
                raise ValueError(f"#{i} substituir: '{v['de'][:60]}' não encontrado.")
            print(f"#{i} substituir: {total} ocorrência(s)")
        elif op == "substituir_paragrafo":
            definir_texto(unico(pars, v["contendo"]), v["texto"])
            print(f"#{i} substituir_paragrafo: ok")
        elif op == "inserir_apos":
            ancora = unico(pars, v["contendo"])
            novo = copy.deepcopy(ancora._p)
            ancora._p.addnext(novo)
            definir_texto(Paragraph(novo, ancora._parent), v["texto"])
            print(f"#{i} inserir_apos: ok")
        elif op == "remover_paragrafo":
            p = unico(pars, v["contendo"])
            p._p.getparent().remove(p._p)
            print(f"#{i} remover_paragrafo: ok")
        elif op == "historico":
            tabela = next((t for t in d.tables if t.rows and t.rows[0].cells[0].text.strip().lower() == "versão"), None)
            if tabela is None:
                raise ValueError(f"#{i} historico: tabela 'Versão' não encontrada.")
            nova = copy.deepcopy(tabela.rows[-1]._tr)
            tabela.rows[-1]._tr.addnext(nova)
            linha = Table(tabela._tbl, tabela._parent).rows[-1]
            for cel, val in zip(linha.cells, [v.get("versao", ""), v.get("data", ""), v.get("autor", ""), v.get("descricao", "")]):
                if cel.paragraphs:
                    definir_texto(cel.paragraphs[0], str(val))
                    for extra in cel.paragraphs[1:]:
                        extra._p.getparent().remove(extra._p)
            print(f"#{i} historico: linha {v.get('versao')} adicionada")
        else:
            raise ValueError(f"#{i}: operação desconhecida '{op}'")
    saida.parent.mkdir(parents=True, exist_ok=True)
    d.save(saida)
    for a in avisos:
        print("[aviso]", a)
    print(f"Gerado: {spec['saida']}")
    print(f"Próximo passo: powershell -ExecutionPolicy Bypass -File scripts/finalizar_docx.ps1 \"{spec['saida']}\"")


def listar(caminho: Path, filtro: str | None) -> None:
    d = docx.Document(caminho)
    for n, p in enumerate(todos_paragrafos(d)):
        if p.text.strip() and (not filtro or filtro.lower() in p.text.lower()):
            estilo = p.style.name if p.style is not None else ""
            print(f"[{n}] ({estilo}) {p.text[:200]}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("arquivo", type=Path, help="especificação .yaml, ou .docx com --listar")
    ap.add_argument("--listar", action="store_true", help="lista os parágrafos do .docx (para localizar trechos)")
    ap.add_argument("--filtro", help="com --listar: mostra só parágrafos que contêm o texto")
    a = ap.parse_args()
    if a.listar:
        listar(a.arquivo, a.filtro)
        return
    spec = yaml.safe_load(a.arquivo.read_text(encoding="utf-8"))
    try:
        aplicar(spec)
    except ValueError as e:
        sys.exit(f"Nenhuma alteração gravada. {e}")


if __name__ == "__main__":
    main()
