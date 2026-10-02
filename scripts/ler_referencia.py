"""Converte arquivos de referência OPCIONAIS em texto para análise pelo Claude.

Uso:
    uv run python scripts/ler_referencia.py                 # todos os arquivos de projeto/referencias/
    uv run python scripts/ler_referencia.py <arquivo> [--abas Aba1,Aba2]

Gera projeto/referencias/_texto/<arquivo>.md (não altera o original).
Formatos: .pptx (texto por slide, tabelas e marcas de formatação relevantes),
.xlsx (abas como tabelas), .docx (parágrafos e tabelas), .pdf (texto via pypdf,
se instalado), .md/.txt/.csv (cópia).
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
REFS = RAIZ / "projeto" / "referencias"
SAIDA = REFS / "_texto"
MAX_LINHAS_ABA = 2000


def pptx_md(p: Path) -> str:
    from pptx import Presentation
    from pptx.util import Pt  # noqa: F401

    out = []
    for n, slide in enumerate(Presentation(p).slides, 1):
        out.append(f"\n## Slide {n}\n")
        for shp in slide.shapes:
            if shp.has_text_frame:
                for par in shp.text_frame.paragraphs:
                    partes = []
                    for r in par.runs:
                        t = r.text
                        if not t.strip():
                            partes.append(t)
                            continue
                        if r.font.bold:
                            t = f"**{t.strip()}**"
                        if r.font.strike if hasattr(r.font, "strike") else False:
                            t = f"~~{t}~~"
                        partes.append(t)
                    linha = "".join(partes).strip()
                    if linha:
                        out.append(f"- {linha}")
            if getattr(shp, "has_table", False) and shp.has_table:
                linhas = [[c.text.strip() for c in row.cells] for row in shp.table.rows]
                if linhas:
                    out.append("")
                    out.append("| " + " | ".join(linhas[0]) + " |")
                    out.append("|" + "---|" * len(linhas[0]))
                    out += ["| " + " | ".join(l) + " |" for l in linhas[1:]]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(f"\n> Notas: {slide.notes_slide.notes_text_frame.text.strip()}")
    out.append("\n> Observação: caracteres sobrescritos (¹ ² ³) e negrito podem indicar vínculos/destaques; confirme o significado com o usuário.")
    return "\n".join(out)


def xlsx_md(p: Path, abas: list[str] | None) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    out = [f"Abas disponíveis: {', '.join(wb.sheetnames)}"]
    for ws in wb.worksheets:
        if abas and ws.title not in abas:
            out.append(f"\n## Aba '{ws.title}' (ignorada por filtro)")
            continue
        out.append(f"\n## Aba '{ws.title}'\n")
        n = 0
        for row in ws.iter_rows(values_only=True):
            if not any(c is not None for c in row):
                continue
            vals = ["" if c is None else str(c).replace("\n", " ⏎ ").replace("|", "/") for c in row]
            out.append("| " + " | ".join(vals) + " |")
            if n == 0:
                out.append("|" + "---|" * len(vals))
            n += 1
            if n >= MAX_LINHAS_ABA:
                out.append(f"(truncado em {MAX_LINHAS_ABA} linhas)")
                break
    return "\n".join(out)


def docx_md(p: Path) -> str:
    import docx

    d = docx.Document(p)
    out = []
    for par in d.paragraphs:
        t = par.text.strip()
        if not t:
            continue
        estilo = (par.style.name or "").lower()
        if estilo.startswith(("heading", "título", "titulo")):
            out.append(f"\n### {t}")
        else:
            out.append(t)
    for i, tb in enumerate(d.tables, 1):
        out.append(f"\n**Tabela {i}**\n")
        for j, row in enumerate(tb.rows):
            vals = [c.text.strip().replace("\n", " ") for c in row.cells]
            out.append("| " + " | ".join(vals) + " |")
            if j == 0:
                out.append("|" + "---|" * len(vals))
    return "\n".join(out)


def pdf_md(p: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "(pypdf não instalado: rode `uv run --with pypdf python scripts/ler_referencia.py <arquivo>`)"
    return "\n\n".join(f"## Página {i}\n{pg.extract_text() or ''}" for i, pg in enumerate(PdfReader(p).pages, 1))


def converter(p: Path, abas: list[str] | None) -> Path:
    ext = p.suffix.lower()
    if ext == ".pptx":
        txt = pptx_md(p)
    elif ext in {".xlsx", ".xlsm"}:
        txt = xlsx_md(p, abas)
    elif ext == ".docx":
        txt = docx_md(p)
    elif ext == ".pdf":
        txt = pdf_md(p)
    elif ext in {".md", ".txt", ".csv"}:
        txt = p.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"formato não suportado: {p.name}")
    SAIDA.mkdir(parents=True, exist_ok=True)
    destino = SAIDA / f"{p.name}.md"
    destino.write_text(f"# {p.name}\n\n{txt}\n", encoding="utf-8")
    return destino


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("arquivo", nargs="?", type=Path)
    ap.add_argument("--abas", help="(xlsx) lista de abas a considerar, separadas por vírgula")
    a = ap.parse_args()
    abas = [x.strip() for x in a.abas.split(",")] if a.abas else None
    arquivos = [a.arquivo] if a.arquivo else sorted(
        f for f in REFS.glob("*") if f.is_file() and not f.name.startswith(("~$", ".")))
    if not arquivos:
        print("Nenhum arquivo de referência em projeto/referencias/ (as referências são opcionais).")
        return
    for f in arquivos:
        try:
            print(f"{f.name} -> {converter(f, abas).relative_to(RAIZ)}")
        except Exception as e:  # noqa: BLE001
            print(f"{f.name}: não convertido ({e})", file=sys.stderr)


if __name__ == "__main__":
    main()
